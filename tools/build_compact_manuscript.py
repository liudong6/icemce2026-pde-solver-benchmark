"""Build the compact Word/LaTeX main paper from one editable content source."""
from pathlib import Path
import json,re
from datetime import datetime,timezone
from docx import Document
from docx.shared import Cm,Pt,Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls,qn
import build_iop_docx as h

ROOT=h.ROOT
TITLE='A Coefficient-Aware Finite-Difference Benchmark for Solver Selection and CPU/GPU Stencil Scaling in Heat-Conduction Simulation'
FACE_CAP='Table 1. All contrast-100 historical face checks: arithmetic/harmonic iterations and setup-inclusive winners (J = Jacobi-PCG; SA = smoothed-aggregation PCG). Single timing passes; no uncertainty estimate.'
CALL_CAP='Table 2. All callback-calibration cells. Winners: method, median total ms and wins/7. Last columns: paired median monitored/count-only solve-time ratios.'
PORT_CAP='Table 3. All 12 interleaved translated pairs. RS = classical AMG-PCG; SA also competed. Stable: each position retains its winner in all seven repeats. E(P) uses medians even for unstable pairs; it is not a confidence bound.'
FIG_CAP='Figure 1. Matched motifs (left); two-island source/alignment sensitivity at N = 192, c = 1000, harmonic faces (right). Bars: median setup + solve time; whiskers: five-run range; labels: iterations.'

def equations(name):
    if name=='model':
        return [(r'-\nabla\!\cdot(k\nabla u)=f,\quad \Omega=(0,1)^2,\quad u|_{\partial\Omega}=0,\qquad (Au)_i=h^{-2}\!\sum_{j\sim i}k_{ij}(u_i-u_j).',
         h.math_seq('−∇·(k∇u) = f,  Ω = ',h.math_sup('(0,1)','2'),',  ',h.math_sub('u|','∂Ω'),' = 0,    ',h.math_sub('(Au)','i'),' = ',h.math_sup('h','−2'),h.math_sub('Σ','j∼i'),h.math_sub('k','ij'),'(',h.math_sub('u','i'),' − ',h.math_sub('u','j'),').'))]
    if name=='faces':
        return [(r'k_H=\frac{2ab}{a+b},\qquad k_A=\frac{a+b}{2},\qquad k_A-k_H=\frac{(a-b)^2}{2(a+b)}.',
         h.math_seq(h.math_sub('k','H'),' = ',h.math_frac('2ab','a+b'),',    ',h.math_sub('k','A'),' = ',h.math_frac('a+b','2'),',    ',h.math_sub('k','A'),' − ',h.math_sub('k','H'),' = ',h.math_frac(h.math_sup('(a−b)','2'),'2(a+b)'),'.'))]
    if name=='floor':
        return [(r'E(P)=\min_{s\in P}\left[\frac12\sum_{j=0}^1\frac{T(j,s)}{\min_{t\in P}T(j,t)}\right]-1.',h.equation_omml_lines('',9)[0])]
    raise ValueError(name)

def escape_tex(text,keys):
    citations=[];protected=[]
    def protect(value):
        protected.append(value);return f'TEXSAFEPLACEHOLDER{len(protected)-1}END'
    text=re.sub(r'https?://[^\s]+',lambda m:protect(r'\url{'+m[0]+'}'),text)
    text=text.replace('D^(−1/2)',protect(r'\ensuremath{D^{-1/2}}'))
    text=re.sub(r'\b(\d+(?:\.\d+)?)e([−-]?\d+)',lambda m:protect(r'\ensuremath{'+m[1]+r'\times10^{'+m[2].replace('−','-')+'}}'),text)
    def cite(m):
        citations.append(r'\cite{'+','.join(keys[int(n)-1] for n in m[1].split(','))+'}')
        return f'CITATIONPLACEHOLDER{len(citations)-1}END'
    text=re.sub(r'\[\[([^\]]+)\]\]',cite,text)
    text=re.sub(r'([%&#_$])',r'\\\1',text)
    text=text.replace('^',r'\textasciicircum{}')
    replacements={'−':r'\ensuremath{-}','–':'--','→':r'\ensuremath{\to}','≤':r'\ensuremath{\leq}','×':r'\ensuremath{\times}','π':r'\ensuremath{\pi}','²':r'\ensuremath{^{2}}','√':r'\ensuremath{\sqrt{2}}','’':"'"}
    text=text.replace('√2',r'\ensuremath{\sqrt{2}}')
    for a,b in replacements.items():text=text.replace(a,b)
    for i,c in enumerate(citations):text=text.replace(f'CITATIONPLACEHOLDER{i}END',c)
    for i,c in enumerate(protected):text=text.replace(f'TEXSAFEPLACEHOLDER{i}END',c)
    text=text.replace('centred/transposed/x-shifted',r'centred/\allowbreak{}transposed/\allowbreak{}x-shifted')
    return text

def paragraph(doc,text,style):
    p=doc.add_paragraph(style=style)
    pattern=r'(D\^\(−1/2\)|\b\d+(?:\.\d+)?e[−-]?\d+)'
    for token in re.split(pattern,text):
        if token=='D^(−1/2)':math=h.math_sup('D','−1/2')
        elif re.fullmatch(r'\d+(?:\.\d+)?e[−-]?\d+',token):
            coeff,exp=token.split('e');math=h.math_seq(coeff,' × ',h.math_sup('10',exp))
        else:h.add_runs_with_citations(p,token);continue
        p._p.append(parse_xml(f'<m:oMath {nsdecls("m")}>{math}</m:oMath>'))
    h.normalize_runs(p);return p

def compact_table(doc,name):
    if name=='portfolio':
        h.add_data_table(doc,PORT_CAP,['N','Motif','Source','Centred','Shifted','Stable','E(P), %'],h.portfolio_table_rows())
        return 'portfolio'
    headers=['Case','N','CG','J','SA','Winner'] if name=='face' else ['Case','N','Monitored','Count-only','CG','J','SA']
    filename='face_table_rows.csv' if name=='face' else 'table_rows.csv'
    rows=[list(r.values()) for r in h.read_csv_rows(ROOT/'results/analysis/protocol_sensitivity'/filename)]
    h.add_data_table(doc,FACE_CAP if name=='face' else CALL_CAP,headers,rows)
    table=doc.tables[-1]
    table.autofit=False
    widths=[2.42,1.0,2.8,2.65,2.65,2.7] if name=='face' else [2.5,1.1,3.5,3.5,1.1,1.1,1.1]
    for col,w in zip(table.columns,widths):col.width=Cm(w)
    for row in table.rows:
        for cell,w in zip(row.cells,widths):
            cell.width=Cm(w)
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
    return filename

def build_compact():
    metadata=json.loads((ROOT/'paper/compact_references.json').read_text(encoding='utf-8'))
    keys=metadata['keys'];refs=metadata['entries']
    source=(ROOT/'paper/compact_manuscript.txt').read_text(encoding='utf-8')
    chunks=source.strip().split('\n\n')
    doc=Document(h.TEMPLATE);h.clear_body(doc);h.set_document_defaults(doc)
    # Preserve the template's inherited 11 pt body rather than the old builder's 10 pt fallback.
    for sty in ['Normal','IOP-CS-BodyText','IOP-CS-BodyNoIndent','IOP-CS-SectionHead','IOP-CS-SubsectionHeading']:
        doc.styles[sty].font.size=Pt(11)
    doc.core_properties.author='Dong Liu; Yuanxin Wang';doc.core_properties.last_modified_by='Dong Liu';doc.core_properties.title=TITLE
    doc.core_properties.modified=datetime.now(timezone.utc)
    h.add_paragraph(doc,TITLE,'IOP-CS-Title',WD_ALIGN_PARAGRAPH.CENTER)
    h.add_paragraph(doc,'Dong Liu* and Yuanxin Wang','IOP-CS-Author',WD_ALIGN_PARAGRAPH.CENTER)
    h.add_paragraph(doc,'University of Nottingham Ningbo China, Ningbo 315100, China','IOP-CS-Affiliation',WD_ALIGN_PARAGRAPH.CENTER)
    h.add_paragraph(doc,'*E-mail: ssydl3@nottingham.edu.cn','IOP-CS-Affiliation',WD_ALIGN_PARAGRAPH.CENTER)
    tex=[r'\documentclass[10pt]{article}',r'\usepackage[a4paper,top=4cm,bottom=2.7cm,left=2.5cm,right=2.5cm]{geometry}',r'\usepackage[T1]{fontenc}',r'\usepackage[utf8]{inputenc}',r'\usepackage{amsmath,amssymb,booktabs,graphicx,url}',r'\usepackage[hidelinks]{hyperref}',r'\setlength{\emergencystretch}{2em}',r'\title{'+TITLE+'}',r'\author{Dong Liu* and Yuanxin Wang\\University of Nottingham Ningbo China\\Ningbo, China\\\texttt{ssydl3@nottingham.edu.cn}}',r'\date{}',r'\begin{document}',r'\maketitle']
    eqnum=0;abstract=False;ack=False
    for block in chunks:
        if block.startswith('::abstract\n'):
            text=block.split('\n',1)[1]
            p=h.add_paragraph(doc,'Abstract. '+text,'IOP-CS-Abstract')
            # IOP requires 10 mm after the abstract.
            p.paragraph_format.space_after=Cm(1.0)
            p.paragraph_format.space_before=Pt(8)
            tex += [r'\begin{abstract}',escape_tex(text,keys),r'\end{abstract}']
        elif block.startswith('::section '):
            section_parts=block.removeprefix('::section ').split('\n',1)
            head=section_parts[0];number,title=head.split('. ',1)
            h.add_section(doc,int(number),title);tex.append(r'\section{'+title+'}')
            if len(section_parts)>1:
                text=section_parts[1]
                paragraph(doc,text,'IOP-CS-BodyNoIndent');tex.append(escape_tex(text,keys)+r'\par')
        elif block.startswith('::equation '):
            for latex,omml in equations(block.split()[1]):
                eqnum+=1
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.keep_together=True
                p._p.append(parse_xml(f'<m:oMath {nsdecls("m")}>{omml}</m:oMath>'))
                p.add_run(f'    ({eqnum})').font.size=Pt(10)
                tex += [r'\begin{equation}',latex,r'\end{equation}']
        elif block.startswith('::table '):
            name=block.split()[1];compact_table(doc,name)
            # Source tables are unchanged; the compact main uses shorter captions.
            fn={'face':'face_sensitivity_main','callback':'protocol_sensitivity','portfolio':'portfolio_summary'}[name]
            raw=(ROOT/'paper/tables'/f'{fn}.tex').read_text(encoding='utf-8')
            cap={'face':FACE_CAP,'callback':CALL_CAP,'portfolio':PORT_CAP}[name]
            raw=re.sub(r'\\caption\{[^\n]+\}',lambda m:r'\caption{'+escape_tex(cap.split('. ',1)[1],keys)+'}',raw,count=1)
            tex.append(raw)
        elif block.startswith('::figure '):
            h.add_figure(doc,'compact_counterfactual.png',FIG_CAP,width_inches=6.15)
            tex += [r'\begin{figure}[htbp]\centering',r'\includegraphics[width=\linewidth]{figures/compact_counterfactual.png}',r'\caption{'+escape_tex(FIG_CAP.split('. ',1)[1],keys)+r'}\end{figure}']
        elif block.startswith('::contributions\n'):
            p=doc.add_paragraph('Author contributions','IOP-CS-SectionHead');p.paragraph_format.keep_with_next=True
            text=block.split('\n',1)[1]
            paragraph(doc,text,'IOP-CS-BodyNoIndent')
            tex += [r'\section*{Author contributions}',escape_tex(text,keys)+r'\par']
        elif block.startswith('::acknowledgements\n'):
            p=doc.add_paragraph('Acknowledgements','IOP-CS-SectionHead');p.paragraph_format.keep_with_next=True
            text=block.split('\n',1)[1];h.add_paragraph(doc,text,'IOP-CS-BodyNoIndent')
            tex += [r'\section*{Acknowledgements}',escape_tex(text,keys)]
        else:
            paragraph(doc,block,'IOP-CS-BodyText');tex.append(escape_tex(block,keys)+r'\par')
    for p in doc.paragraphs:
        for name in ['Dong Liu','Yuanxin Wang']:
            if p.text.startswith(name+' ') and p.runs and p.runs[0].text.startswith(name):
                run=p.runs[0];rest=run.text[len(name):];run.text=name;run.bold=True
                new=p.add_run(rest);run._r.addnext(new._r)
    p=doc.add_paragraph('References','IOP-CS-SectionHead');p.paragraph_format.keep_with_next=True
    tex += [r'\begin{thebibliography}{99}',r'\small\setlength{\itemsep}{0pt}\setlength{\parskip}{0pt}']
    for i,(key,ref) in enumerate(zip(keys,refs),1):
        p=doc.add_paragraph(style='IOP-CS-ReferenceText');p.add_run(f'[{i}] {ref}');h.normalize_runs(p)
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
        for run in p.runs:run.font.size=Pt(9)
        tex.append(r'\bibitem{'+key+'} '+escape_tex(ref,keys))
    tex += [r'\end{thebibliography}',r'\end{document}']
    (ROOT/'paper/main.tex').write_text('\n'.join(tex)+'\n',encoding='utf-8')
    doc.save(h.OUT);h.normalize_docx_revision(h.OUT)
    from highlight_manuscript_changes import highlight
    highlight(ROOT/'tools/review_baseline.docx',h.OUT,h.OUT,ROOT/'paper/REVISION_HIGHLIGHTS.json')
    print(h.OUT)

if __name__=='__main__':build_compact()
