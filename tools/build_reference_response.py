"""Build the response to the reference-focused review from the audited counts."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
audit=json.loads((ROOT/'paper/REFERENCE_AUDIT.json').read_text(encoding='utf-8'))
c=audit['counts']
assert c['total']==31 and c['recent_2022_2026']==18 and c['preprints']==1
title='A Coefficient-Aware Finite-Difference Benchmark for Solver Selection and CPU/GPU Stencil Scaling in Heat-Conduction Simulation'
request='Revise the references so that more than half are from 2022–2026, use foreign-language literature wherever possible, and keep preprints at no more than 20% of the total. Highlight the manuscript changes in yellow.'
response=[
'Thank you for pointing out the need to update the literature coverage. We have revised the reference list and the related-work discussion, with particular attention to recent work on solver selection, AMG coarsening and sparse computation.',
f'The revised manuscript contains {c["total"]} references. Of these, {c["recent_2022_2026"]} ({c["recent_percent"]:.2f}%) are dated 2022–2026. Even when the retained preprint is excluded from the recent-publication count, {c["recent_formal_papers"]} formally published journal or conference papers fall in this period ({c["recent_formal_percent"]:.2f}% of all references). All 31 entries are English-language sources. One reference is a preprint ({c["preprint_percent"]:.2f}%), below the 20% limit.',
'We added eight recent formal publications and explain their relevance in Section 1: matching-based AMG coarsening [16], the Ginkgo framework [17], mixed-precision AMG and sparse matrix-vector products [20,21], and feature-based, multimodal and online solver selection [26–29]. The revised discussion distinguishes these methods from our matched-field audit and makes clear that they were not additional algorithms evaluated in our experiments.',
'We also replaced the 2009 arXiv citation by Aksoylu and Yeter with its published journal version [23]. The material-aware AMG study [25] remains explicitly identified as a preprint. We removed four general or overlapping entries while retaining the foundational references needed for the discretisation, Krylov and multigrid methods, interface treatment and timing methodology. Publication years follow the cited editions or journal volumes; web access dates are not counted as publication years.',
'Changes to the manuscript text, reference entries and citation numbering are highlighted in yellow relative to the version reviewed in this round. The numerical results and conclusions are unchanged.',
]
d=Document();sec=d.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.85)
for name in ['Normal','Title','Heading 1']:
 s=d.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
 for f in s.element.xpath('.//w:rFonts'):
  for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:f.attrib.pop(qn('w:'+a),None)
for s in d.styles:
 for border in s.element.xpath('.//w:pBdr'):border.getparent().remove(border)
d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.line_spacing=1.08;d.styles['Normal'].paragraph_format.space_after=Pt(9)
d.styles['Title'].font.size=Pt(16)
d.core_properties.title='Response to the reference revision request';d.core_properties.author='Dong Liu';d.core_properties.subject=title
d.add_paragraph('Response to the reference revision request','Title')
p=d.add_paragraph();p.add_run('Manuscript: ').bold=True;p.add_run(title)
d.add_paragraph('Dear Editor and Reviewer,')
p=d.add_paragraph();p.add_run('Reviewer request (translated and summarised)').bold=True
p=d.add_paragraph(request)
for r in p.runs:r.italic=True
p=d.add_paragraph();p.add_run('Response').bold=True
for t in response:d.add_paragraph(t)
p=d.add_paragraph();p.add_run('Revised locations: ').bold=True;p.add_run('Section 1, citation numbering throughout the manuscript, and the reference list. The six-page main paper retains the review evidence; detailed methods and the complete original result displays accompany it as a supplement.')
d.add_paragraph('Sincerely,\nDong Liu')
d.save(ROOT/'paper/RESPONSE_TO_REVIEWERS.docx')
(ROOT/'paper/RESPONSE_TO_REVIEWERS.md').write_text('\n\n'.join(p.text for p in d.paragraphs if p.text)+'\n',encoding='utf-8')
print(ROOT/'paper/RESPONSE_TO_REVIEWERS.docx')
